import streamlit as st
import csv
import gzip
import io
import json
import math
import re
import statistics
import urllib3
from datetime import date, timedelta, datetime, time
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import pandas as pd
import folium
from openpyxl.utils import get_column_letter
from shapely.geometry import Point
from shapely.geometry import shape as shapely_shape
from streamlit_folium import st_folium
from zoneinfo import ZoneInfo
from weather_sources import build_weather_dataset

# Suppress SSL warnings
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# === Metric Configuration ===
PARAMETERS = {
    "Temperature @2m (°C)": ("T2M", True),
    "Relative Humidity @2m (%)": ("RH2M", True),
    "Wind Speed @2m (m/s)": ("WS2M", True),
    "Wind Direction @2m (°)": ("WD2M", True),
    "Precipitation (mm/day) [Daily Only]": ("PRECTOTCORR", False), 
}

# === HEADER RENAMING MAP ===
HEADER_MAP = {
    "T2M": "TEMP_C",
    "RH2M": "RELHUM_%",
    "WS2M": "WS_MPS",
    "WD2M": "WD_DEGREES",
    "WD2M_COMPASS": "WD_CARDINAL",
    "PRECTOTCORR": "PREC_MM"
}

DEFAULT_COMMUNITY = "AG"
DEFAULT_TIME_STANDARD = "LST"
SOURCE_STRATEGIES = ["Auto", "NASA only", "Prefer INMET"]
TIMEZONE_DATA_PATH = Path(__file__).resolve().parent / "data" / "timezones_americas.geojson.gz"


@st.cache_resource(show_spinner=False)
def _load_timezone_polygons() -> List[Tuple[str, object]]:
    with gzip.open(TIMEZONE_DATA_PATH, "rt", encoding="utf-8") as fh:
        data = json.load(fh)
    return [(feat["properties"]["tzid"], shapely_shape(feat["geometry"])) for feat in data["features"]]


def timezone_at(lat: float, lon: float) -> Optional[str]:
    point = Point(lon, lat)
    for tzid, geom in _load_timezone_polygons():
        if geom.contains(point):
            return tzid
    return None


def tr(en: str, es: str, pt: Optional[str] = None) -> str:
    lang = st.session_state.get("ui_language", "en")
    if lang == "es":
        return es
    if lang == "pt":
        return pt if pt is not None else en
    return en


def _normalize_headers() -> Dict[str, str]:
    ctx = getattr(st, "context", None)
    if ctx is None:
        return {}
    try:
        return {str(k).lower(): str(v) for k, v in dict(ctx.headers).items()}
    except Exception:
        return {}


def _detect_runtime_defaults() -> Dict[str, object]:
    headers = _normalize_headers()
    accept_language = headers.get("accept-language", "").lower()
    if re.search(r"(^|,|\s)pt(?:-|;|,|$)", accept_language):
        language = "pt"
        source = "accept_language"
    elif re.search(r"(^|,|\s)es(?:-|;|,|$)", accept_language):
        language = "es"
        source = "accept_language"
    else:
        language = "en"
        source = "default"

    return {
        "language": language,
        "language_source": source,
    }


@st.cache_data(ttl=86400, show_spinner=False)
def infer_utc_offset_from_coordinates(lat: float, lon: float, ref_date_iso: str) -> Optional[int]:
    try:
        tz_name = timezone_at(lat, lon)
        if not tz_name:
            return None
        ref_date = date.fromisoformat(ref_date_iso)
        ref_dt = datetime(ref_date.year, ref_date.month, ref_date.day, 12, 0, tzinfo=ZoneInfo(tz_name))
        offset = ref_dt.utcoffset()
        if offset is None:
            return None
        return int(round(offset.total_seconds() / 3600.0))
    except Exception:
        return None

# --- Utilities ---
def deg_to_compass_16(deg):
    try:
        d = float(deg) % 360.0
    except (TypeError, ValueError):
        return ""
    dirs = ["N","NNE","NE","ENE","E","ESE","SE","SSE","S","SSW","SW","WSW","W","WNW","NW","NNW"]
    return dirs[int((d + 11.25) // 22.5) % 16]

def vector_average_degrees(angles):
    if not angles: return None
    sin_sum = sum(math.sin(math.radians(a)) for a in angles)
    cos_sum = sum(math.cos(math.radians(a)) for a in angles)
    avg_rad = math.atan2(sin_sum / len(angles), cos_sum / len(angles))
    avg_deg = math.degrees(avg_rad)
    return avg_deg % 360.0

def _normalize_decimal_separator(value: str) -> str:
    # Some Windows locales default to ',' as the decimal separator; accept it
    # transparently alongside '.'.
    return str(value or "").strip().replace(",", ".")

def valid_lat_lon(lat_str, lon_str):
    try:
        lat = float(_normalize_decimal_separator(lat_str))
        lon = float(_normalize_decimal_separator(lon_str))
        if not (-90.0 <= lat <= 90.0 and -180.0 <= lon <= 180.0):
            return None, None
        return lat, lon
    except ValueError:
        return None, None

def dms_to_decimal(dms_str, is_lat=True):
    txt = str(dms_str or "").strip().upper()
    if not txt:
        return None

    # Normalize a ',' used as a decimal separator within a number (e.g. the
    # seconds component) to '.' before extracting numeric tokens.
    txt = re.sub(r"(\d),(\d)", r"\1.\2", txt)

    hemi_match = re.search(r"[NSEW]", txt)
    hemi = hemi_match.group(0) if hemi_match else None
    nums = re.findall(r"[-+]?\d+(?:\.\d+)?", txt)
    if len(nums) == 0:
        return None

    try:
        deg = float(nums[0])
        minutes = float(nums[1]) if len(nums) > 1 else 0.0
        seconds = float(nums[2]) if len(nums) > 2 else 0.0
    except ValueError:
        return None

    if minutes < 0 or minutes >= 60 or seconds < 0 or seconds >= 60:
        return None

    value = abs(deg) + (minutes / 60.0) + (seconds / 3600.0)
    sign = -1 if deg < 0 else 1

    if hemi in ("S", "W"):
        sign = -1
    elif hemi in ("N", "E"):
        sign = 1

    value *= sign
    if is_lat and not (-90.0 <= value <= 90.0):
        return None
    if (not is_lat) and not (-180.0 <= value <= 180.0):
        return None
    return value

def build_coordinate_preview_map(lat: float, lon: float) -> folium.Map:
    m = folium.Map(location=[lat, lon], zoom_start=13, tiles=None, control_scale=True)
    folium.TileLayer(
        tiles="https://server.arcgisonline.com/ArcGIS/rest/services/World_Imagery/MapServer/tile/{z}/{y}/{x}",
        attr="Esri, Maxar, Earthstar Geographics",
        name="Satellite",
        overlay=False,
        control=False,
    ).add_to(m)
    folium.TileLayer(
        tiles="https://server.arcgisonline.com/ArcGIS/rest/services/Reference/World_Boundaries_and_Places/MapServer/tile/{z}/{y}/{x}",
        attr="Esri",
        name="Labels",
        overlay=True,
        control=False,
    ).add_to(m)
    # Use an icon-font marker (not Leaflet's default PNG marker) since the default
    # relies on auto-detecting a relative image path, which is unreliable inside
    # the sandboxed iframe st_folium renders the map in.
    folium.Marker(
        location=[lat, lon],
        icon=folium.Icon(color="red", icon="map-pin", prefix="fa"),
    ).add_to(m)
    return m

def get_precip_sum(start_d, end_d, data_dict):
    total = 0.0
    cur_d = start_d
    while cur_d <= end_d:
        dt_str = cur_d.strftime("%Y%m%d")
        val = data_dict.get(dt_str, {}).get("PRECTOTCORR", 0.0)
        if val != -999.0 and val is not None:
            total += val
        cur_d += timedelta(days=1)
    return round(total, 2)

def to_arm_date(date_key):
    d = datetime.strptime(date_key, "%Y%m%d")
    return d.strftime("%d-%b-%y").lstrip("0")

# --- Streamlit UI Setup ---
st.set_page_config(page_title="Weather2ARM", layout="wide", page_icon="🌦️")

# Custom CSS for Bottom-Right NASA Rights
st.markdown(
    """
    <style>
    .nasa-footer {
        position: fixed;
        bottom: 10px;
        right: 15px;
        font-size: 11px;
        color: #888888;
        background-color: rgba(255, 255, 255, 0.8);
        padding: 5px 10px;
        border-radius: 5px;
        text-align: right;
        z-index: 100;
        max-width: 350px;
        pointer-events: none;
    }
    @media (prefers-color-scheme: dark) {
        .nasa-footer {
            background-color: rgba(14, 17, 23, 0.8);
            color: #aaaaaa;
        }
    }
    </style>
    <div class="nasa-footer">
        <b>NASA POWER + INMET Data References</b><br>
        These data were obtained from the NASA Langley Research Center (LaRC) POWER Project funded through the NASA Earth Science/Applied Science Program.
        <br><br>
        INMET station data are provided by Instituto Nacional de Meteorologia (INMET), Brazil.
    </div>
    """,
    unsafe_allow_html=True
)

# Initialize Session State Variables
if "csv_hourly_str" not in st.session_state: st.session_state.csv_hourly_str = None
if "csv_daily_str" not in st.session_state: st.session_state.csv_daily_str = None
if "excel_hourly_arm" not in st.session_state: st.session_state.excel_hourly_arm = None
if "excel_daily_arm" not in st.session_state: st.session_state.excel_daily_arm = None
if "excel_app_format" not in st.session_state: st.session_state.excel_app_format = None
if "output_metadata_json" not in st.session_state: st.session_state.output_metadata_json = None
if "base_filename" not in st.session_state: st.session_state.base_filename = ""
if "is_arm" not in st.session_state: st.session_state.is_arm = False
if "auto_defaults_ready" not in st.session_state:
    defaults = _detect_runtime_defaults()
    st.session_state.ui_language = defaults.get("language", "en")
    st.session_state.language_detection_source = defaults.get("language_source", "default")
    st.session_state.ui_language_user_edited = False
    st.session_state.local_utc_offset = -3
    st.session_state.detected_utc_offset = None
    st.session_state.utc_offset_source = "default"
    st.session_state.utc_offset_user_edited = False
    st.session_state.last_autodetected_offset = None
    st.session_state.auto_defaults_ready = True
if "ui_language" not in st.session_state: st.session_state.ui_language = "en"
if "ui_language_user_edited" not in st.session_state: st.session_state.ui_language_user_edited = False
if "local_utc_offset" not in st.session_state: st.session_state.local_utc_offset = -3
if "detected_utc_offset" not in st.session_state: st.session_state.detected_utc_offset = None
if "utc_offset_source" not in st.session_state: st.session_state.utc_offset_source = "default"
if "utc_offset_user_edited" not in st.session_state: st.session_state.utc_offset_user_edited = False
if "last_autodetected_offset" not in st.session_state: st.session_state.last_autodetected_offset = None


def _mark_utc_offset_user_edited():
    st.session_state.utc_offset_user_edited = True


def _mark_language_user_edited():
    st.session_state.ui_language = st.session_state.ui_language_selector
    st.session_state.ui_language_user_edited = True
    st.session_state.language_detection_source = "manual"

# Main Layout
lang_col, geo_col = st.columns([1, 2])
if "ui_language_selector" not in st.session_state:
    st.session_state.ui_language_selector = st.session_state.ui_language
with lang_col:
    st.selectbox(
        tr("Language", "Idioma", "Idioma"),
        options=["en", "es", "pt"],
        index={"en": 0, "es": 1, "pt": 2}.get(st.session_state.ui_language_selector, 0),
        format_func=lambda code: "English" if code == "en" else ("Español" if code == "es" else "Português"),
        key="ui_language_selector",
        on_change=_mark_language_user_edited,
    )
st.session_state.ui_language = st.session_state.ui_language_selector

with geo_col:
    st.caption(
        tr(
            f"Language source: {st.session_state.get('language_detection_source', 'default')} (browser language).",
            f"Origen del idioma: {st.session_state.get('language_detection_source', 'default')} (idioma del navegador).",
            f"Fonte do idioma: {st.session_state.get('language_detection_source', 'default')} (idioma do navegador).",
        )
    )

st.title("🌦️ Weather2ARM")
st.markdown(
    tr(
        "Download and process weather data for ARM software using NASA POWER and INMET sources.",
        "Descarga y procesa datos meteorológicos para ARM usando NASA POWER e INMET.",
        "Baixe e processe dados meteorológicos para software ARM usando fontes NASA POWER e INMET.",
    )
)

# 1. Location
st.subheader(tr("1. Location", "1. Ubicación", "1. Localização"))
coord_mode = st.selectbox(
    tr("Coordinate Input Format", "Formato de coordenadas", "Formato de coordenadas"),
    ["decimal", "dms"],
    index=0,
    format_func=lambda mode: tr("Decimal Degrees", "Grados decimales", "Graus decimais") if mode == "decimal" else tr("GMS (Degrees Minutes Seconds)", "GMS (Grados Minutos Segundos)", "GMS (Graus Minutos Segundos)"),
)
coord_col, map_col = st.columns([1, 1])
with coord_col:
    if coord_mode == "decimal":
        lat_input = st.text_input(tr("Latitude", "Latitud", "Latitude"), value="", placeholder="-26.9386111")
        lon_input = st.text_input(tr("Longitude", "Longitud", "Longitude"), value="", placeholder="-52.39805555")
        lat_dms_input, lon_dms_input = "", ""
        st.caption(
            tr(
                "Decimal format note: '.' or ',' are both accepted as decimal separator (example: -26.9386111 or -26,9386111).",
                "Nota para formato decimal: se aceptan '.' o ',' como separador decimal (ejemplo: -26.9386111 o -26,9386111).",
                "Nota para formato decimal: '.' ou ',' são aceitos como separador decimal (exemplo: -26.9386111 ou -26,9386111).",
            )
        )
    else:
        lat_dms_input = st.text_input(tr("Latitude (GMS)", "Latitud (GMS)", "Latitude (GMS)"), value="", placeholder="26 56 19 S")
        lon_dms_input = st.text_input(tr("Longitude (GMS)", "Longitud (GMS)", "Longitude (GMS)"), value="", placeholder="52 23 53 W")
        lat_input, lon_input = "", ""
        st.caption(
            tr(
                "Accepted examples: 26 56 19 S, 26°56'19\"S, -26 56 19, -51º02'06.17'', 51º02'06,17''W",
                "Ejemplos válidos: 26 56 19 S, 26°56'19\"S, -26 56 19, -51º02'06.17'', 51º02'06,17''W",
                "Exemplos aceitos: 26 56 19 S, 26°56'19\"S, -26 56 19, -51º02'06.17'', 51º02'06,17''W",
            )
        )

preview_lat, preview_lon = (None, None)
if coord_mode == "decimal":
    preview_lat, preview_lon = valid_lat_lon(lat_input, lon_input)
else:
    preview_lat = dms_to_decimal(lat_dms_input, is_lat=True)
    preview_lon = dms_to_decimal(lon_dms_input, is_lat=False)

with map_col:
    if preview_lat is not None and preview_lon is not None:
        st_folium(
            build_coordinate_preview_map(preview_lat, preview_lon),
            key="coord_preview_map",
            returned_objects=[],
            use_container_width=True,
            height=300,
        )
    else:
        st.caption(
            tr(
                "Enter valid coordinates to preview the location on a satellite map.",
                "Ingresa coordenadas válidas para ver la ubicación en un mapa satelital.",
                "Informe coordenadas válidas para visualizar a localização em um mapa de satélite.",
            )
        )

if preview_lat is not None and preview_lon is not None:
    inferred_offset = infer_utc_offset_from_coordinates(preview_lat, preview_lon, date.today().isoformat())
    st.session_state.detected_utc_offset = inferred_offset
    if inferred_offset is not None:
        # Keep user override intact; otherwise update field to inferred local offset.
        if not st.session_state.get("utc_offset_user_edited", False):
            st.session_state.local_utc_offset = inferred_offset
        st.session_state.last_autodetected_offset = inferred_offset
        st.session_state.utc_offset_source = "coordinates"
    else:
        st.session_state.utc_offset_source = "default"

# 2. Date Range
st.subheader(tr("2. Date Range", "2. Rango de fechas", "2. Intervalo de datas"))
today = date.today()
col3, col4 = st.columns(2)
with col3: start_date = st.date_input(tr("Start Date", "Fecha de inicio", "Data inicial"), value=date(today.year, 1, 1))
with col4: end_date = st.date_input(tr("End Date", "Fecha de fin", "Data final"), value=today - timedelta(days=1))

selected_params = {code: True for code, _ in PARAMETERS.values()}

# 3. Output Options
st.subheader(tr("3. Output Options", "3. Opciones de salida", "3. Opções de saída"))
col5, col6 = st.columns(2)
with col5:
    out_daily = st.checkbox(tr("Generate Daily Stats", "Generar estadísticas diarias", "Gerar estatísticas diárias"), value=True)
    out_hourly = st.checkbox(tr("Generate Hourly Data", "Generar datos horarios", "Gerar dados horários"), value=False)
with col6:
    output_format = st.selectbox(
        tr("Output Layout", "Diseño de salida", "Layout de saída"),
        ["csv", "arm"],
        index=1,
        format_func=lambda item: tr("Standard Layout (CSV)", "Formato estándar (CSV)", "Formato padrão (CSV)") if item == "csv" else tr("ARM Software Layout (Excel)", "Formato ARM (Excel)", "Formato ARM (Excel)"),
    )
    apply_precip_filter = st.checkbox(tr("Filter Low Rainfall (Daily)", "Filtrar lluvia baja (diario)", "Filtrar chuva baixa (diário)"), value=True)
    precip_threshold = st.number_input(tr("Rainfall Threshold (mm)", "Umbral de lluvia (mm)", "Limite de chuva (mm)"), value=0.5, step=0.1, disabled=not apply_precip_filter)
st.caption(tr("Rainfall filter applies to NASA POWER daily precipitation values. INMET-primary daily precipitation is not filtered.", "El filtro de lluvia aplica solo a la precipitación diaria de NASA POWER. La precipitación diaria primaria de INMET no se filtra.", "O filtro de chuva se aplica apenas aos valores diários de precipitação da NASA POWER. A precipitação diária primária do INMET não é filtrada."))
st.caption(tr("Hourly precipitation is included in hourly CSV downloads when INMET is the source. NASA POWER provides precipitation at daily resolution only.", "La precipitación horaria se incluye en CSV horario cuando la fuente es INMET. NASA POWER solo entrega precipitación diaria.", "A precipitação horária é incluída nos downloads CSV horários quando a fonte é INMET. A NASA POWER fornece precipitação apenas em resolução diária."))

st.subheader(tr("4. Data Source", "4. Fuente de datos", "4. Fonte de dados"))
col8, col9, col10 = st.columns(3)
with col8:
    source_strategy = st.selectbox(
        tr("Source Selection", "Selección de fuente", "Seleção da fonte"),
        SOURCE_STRATEGIES,
        index=0,
        format_func=lambda item: tr("Auto", "Auto", "Auto") if item == "Auto" else (tr("NASA only", "Solo NASA", "Somente NASA") if item == "NASA only" else tr("Prefer INMET", "Preferir INMET", "Preferir INMET")),
    )
    inmet_gap_fill = st.checkbox(tr("Fill INMET gaps with NASA POWER", "Completar huecos de INMET con NASA POWER", "Preencher lacunas do INMET com NASA POWER"), value=True)
with col9:
    inmet_radius_km = st.number_input(tr("INMET Search Radius (km)", "Radio de búsqueda INMET (km)", "Raio de busca INMET (km)"), min_value=1.0, value=50.0, step=5.0)
    timezone_offset_hours = st.number_input(
        tr("Local UTC Offset (Brasilia default: -3)", "Desfase UTC local (Brasília por defecto: -3)", "Deslocamento UTC local (Brasília padrão: -3)"),
        min_value=-12,
        max_value=14,
        step=1,
        key="local_utc_offset",
        on_change=_mark_utc_offset_user_edited,
    )
    if st.session_state.get("detected_utc_offset") is not None:
        if st.button(tr("Use Auto UTC Offset", "Usar UTC detectado", "Usar UTC automático"), key="use_auto_utc_offset"):
            st.session_state.local_utc_offset = int(st.session_state.detected_utc_offset)
            st.session_state.utc_offset_user_edited = False
            st.rerun()

    detected_utc = st.session_state.get("detected_utc_offset")
    if detected_utc is None:
        st.caption(
            tr(
                "UTC source: default (-3). Enter valid coordinates to auto-detect timezone from location (VPN-safe).",
                "Fuente UTC: valor por defecto (-3). Ingresa coordenadas válidas para detectar la zona horaria por ubicación (compatible con VPN).",
                "Fonte UTC: padrão (-3). Informe coordenadas válidas para detectar automaticamente o fuso horário pela localização (compatível com VPN).",
            )
        )
    else:
        st.caption(
            tr(
                f"UTC source: coordinates -> suggested UTC{int(detected_utc):+d}.",
                f"Fuente UTC: coordenadas -> UTC sugerido {int(detected_utc):+d}.",
                f"Fonte UTC: coordenadas -> UTC sugerido {int(detected_utc):+d}.",
            )
        )
with col10:
    inmet_data_dir = st.text_input(
        tr("INMET Data Directory", "Directorio de datos INMET", "Diretório de dados INMET"),
        value="INMET",
        help=tr(
            "Path relative to the app root. Example: INMET (supports nested folders and ZIP files).",
            "Ruta relativa a la raíz de la app. Ejemplo: INMET (soporta carpetas anidadas y archivos ZIP).",
            "Caminho relativo a raiz do app. Exemplo: INMET (suporta pastas aninhadas e arquivos ZIP).",
        ),
    )
    preferred_inmet_station = st.text_input(
        tr("Preferred INMET Station (optional)", "Estación INMET preferida (opcional)", "Estação INMET preferida (opcional)"),
        value="",
        help=tr("Optional station code or name filter, e.g. A858 or XANXERE.", "Filtro opcional por código o nombre de estación, por ejemplo A858 o XANXERE.", "Filtro opcional por código ou nome da estação, por exemplo A858 ou XANXERE."),
    )
    force_nasa_timezone = st.checkbox(tr("Keep NASA Time Standard Setting", "Mantener configuración de zona horaria NASA", "Manter configuração de fuso horário da NASA"), value=True)
st.info(tr("INMET dataset availability: monthly station files generally cover dates up to the end of the previous month and currently go back to 2025.", "Disponibilidad del dataset INMET: los archivos mensuales suelen cubrir hasta el fin del mes anterior y actualmente llegan hasta 2025.", "Disponibilidade do dataset INMET: os arquivos mensais geralmente cobrem datas até o fim do mês anterior e atualmente retrocedem até 2025."))

# Hidden Menus
with st.expander(tr("🌱 Weather Application Export (ARM Format)", "🌱 Exportación de Aplicaciones (Formato ARM)", "🌱 Exportação de Aplicações (Formato ARM)")):
    st.caption(tr("Generate a secondary structured Excel sheet configured for agronomic software input (Requires Daily Stats).", "Genera una hoja Excel estructurada para entrada de software agronómico (requiere estadísticas diarias).", "Gera uma planilha Excel estruturada para entrada em software agronômico (requer estatísticas diárias)."))
    enable_app_format = st.checkbox(tr("Enable Application Formatting", "Habilitar formato de aplicaciones", "Habilitar formatação de aplicações"), value=False)
    app_dates_input = []
    if enable_app_format:
        num_apps = st.number_input(tr("Number of Applications", "Número de aplicaciones", "Número de aplicações"), min_value=1, max_value=20, value=1)
        for i in range(num_apps):
            app_letter = chr(65 + i)
            d = st.date_input(tr(f"Application {app_letter} Date", f"Fecha de aplicación {app_letter}", f"Data da aplicação {app_letter}"), value=date.today() - timedelta(days=7), key=f"app_{app_letter}")
            t = st.time_input(tr(f"Application {app_letter} Time", f"Hora de aplicación {app_letter}", f"Hora da aplicação {app_letter}"), value=time(9, 0), step=1800, key=f"app_t_{app_letter}")
            app_dates_input.append((app_letter, d, t))
        st.caption(tr("⚠️ Ensure your Date Range covers at least 14 days prior and 28 days after your application dates.", "⚠️ Asegura que el rango de fechas cubra al menos 14 días antes y 28 días después de las fechas de aplicación.", "⚠️ Garanta que o intervalo de datas cubra pelo menos 14 dias antes e 28 dias após as datas de aplicação."))

with st.expander(tr("⚙️ Advanced Settings", "⚙️ Configuración avanzada", "⚙️ Configurações avançadas")):
    ssl_verify = st.checkbox(tr("Enable SSL Verification", "Habilitar verificación SSL", "Habilitar verificação SSL"), value=False)
    debug_mode = st.checkbox(tr("Debug Mode (Show internal logs)", "Modo debug (mostrar logs internos)", "Modo debug (mostrar logs internos)"), value=False)

# --- Processing Engine ---
if st.button(tr("🚀 DOWNLOAD & PROCESS", "🚀 DESCARGAR Y PROCESAR", "🚀 BAIXAR E PROCESSAR"), type="primary", use_container_width=True):
    if coord_mode == "decimal":
        lat, lon = valid_lat_lon(lat_input, lon_input)
    else:
        lat = dms_to_decimal(lat_dms_input, is_lat=True)
        lon = dms_to_decimal(lon_dms_input, is_lat=False)
    if lat is None:
        st.error(tr("❌ Invalid coordinates. Please check your Latitude and Longitude format.", "❌ Coordenadas inválidas. Revisa el formato de latitud y longitud.", "❌ Coordenadas inválidas. Verifique o formato de latitude e longitude."))
        st.stop()
    if start_date > end_date:
        st.error(tr("❌ Start date must be before or equal to End date.", "❌ La fecha de inicio debe ser menor o igual a la fecha de fin.", "❌ A data inicial deve ser menor ou igual a data final."))
        st.stop()
    if not out_hourly and not out_daily:
        st.error(tr("❌ Please select at least one output format (Hourly or Daily).", "❌ Selecciona al menos un formato de salida (horario o diario).", "❌ Selecione pelo menos um formato de saída (horário ou diário)."))
        st.stop()

    user_start_date = start_date
    user_end_date = end_date

    # When application format is enabled, expand fetch range so pre/post windows are complete,
    # including cross-year lookbacks (e.g., Jan application needing previous-year data).
    fetch_start_date = user_start_date
    fetch_end_date = user_end_date
    if enable_app_format and app_dates_input:
        min_app_date = min(d for _, d, _ in app_dates_input)
        max_app_date = max(d for _, d, _ in app_dates_input)
        fetch_start_date = min(user_start_date, min_app_date - timedelta(days=14))
        fetch_end_date = max(user_end_date, max_app_date + timedelta(days=28))

    # Clear previous session state data
    st.session_state.csv_hourly_str = None
    st.session_state.csv_daily_str = None
    st.session_state.excel_hourly_arm = None
    st.session_state.excel_daily_arm = None
    st.session_state.excel_app_format = None
    st.session_state.output_metadata_json = None
    
    st.session_state.is_arm = (output_format == "arm")
    community = DEFAULT_COMMUNITY
    tstd = DEFAULT_TIME_STANDARD
    st.session_state.base_filename = f"POWER_{community}_{lat:.4f}_{lon:.4f}_{user_start_date.strftime('%Y%m%d')}_{user_end_date.strftime('%Y%m%d')}"

    hourly_req = [code for code, sel in selected_params.items() if sel and PARAMETERS[[k for k, v in PARAMETERS.items() if v[0]==code][0]][1]]
    daily_req = [code for code, sel in selected_params.items() if sel and not PARAMETERS[[k for k, v in PARAMETERS.items() if v[0]==code][0]][1]]

    debug_log = []
    daily_storage = {}
    hourly_records = []
    
    with st.status(tr("Fetching and processing weather data...", "Consultando y procesando datos meteorológicos...", "Buscando e processando dados meteorológicos..."), expanded=True) as status:
        weather_result = build_weather_dataset(
            lat=lat,
            lon=lon,
            start_date=fetch_start_date,
            end_date=fetch_end_date,
            selected_params=selected_params,
            community=community,
            tstd=tstd if force_nasa_timezone else "UTC",
            out_daily=out_daily,
            out_hourly=out_hourly,
            enable_app_format=enable_app_format,
            apply_precip_filter=apply_precip_filter,
            precip_threshold=precip_threshold,
            source_strategy=source_strategy,
            inmet_radius_km=float(inmet_radius_km),
            inmet_gap_fill=inmet_gap_fill,
            inmet_data_dir=inmet_data_dir,
            preferred_inmet_station=preferred_inmet_station,
            timezone_offset_hours=int(timezone_offset_hours),
            ssl_verify=ssl_verify,
        )

        daily_storage = weather_result.daily_storage
        hourly_records = weather_result.hourly_records
        output_metadata = weather_result.metadata

        # Keep main outputs constrained to user-selected range while allowing expanded-range
        # data to support application summary windows.
        main_daily_storage = {
            k: v
            for k, v in daily_storage.items()
            if user_start_date.strftime("%Y%m%d") <= str(k) <= user_end_date.strftime("%Y%m%d")
        }
        main_hourly_records = [
            rec
            for rec in hourly_records
            if user_start_date.strftime("%Y%m%d") <= str(rec.get("date_key", "")) <= user_end_date.strftime("%Y%m%d")
        ]

        # Use filtered data for standard outputs.
        daily_storage = main_daily_storage
        hourly_records = main_hourly_records
        st.session_state.output_metadata_json = json.dumps(output_metadata, indent=2, ensure_ascii=False)

        if output_metadata.get("primary_source") == "INMET":
            station_meta = output_metadata.get("station", {})
            st.info(
                tr(
                    f"Using INMET station {station_meta.get('station_code', 'N/A')} - "
                    f"{station_meta.get('station_name', 'Unknown')} "
                    f"({station_meta.get('distance_km', 'N/A')} km).",
                    f"Usando estación INMET {station_meta.get('station_code', 'N/A')} - "
                    f"{station_meta.get('station_name', 'Desconocida')} "
                    f"({station_meta.get('distance_km', 'N/A')} km).",
                    f"Usando estação INMET {station_meta.get('station_code', 'N/A')} - "
                    f"{station_meta.get('station_name', 'Desconhecida')} "
                    f"({station_meta.get('distance_km', 'N/A')} km).",
                )
            )
            if out_hourly and not st.session_state.is_arm:
                st.caption(tr("INMET hourly output includes precipitation (PREC_MM_HR). NASA POWER hourly output does not include precipitation.", "La salida horaria de INMET incluye precipitación (PREC_MM_HR). La salida horaria de NASA POWER no incluye precipitación.", "A saída horária do INMET inclui precipitação (PREC_MM_HR). A saída horária da NASA POWER não inclui precipitação."))
        else:
            st.info(tr(f"Using NASA POWER. Reason: {output_metadata.get('selection_reason', 'N/A')}", f"Usando NASA POWER. Motivo: {output_metadata.get('selection_reason', 'N/A')}", f"Usando NASA POWER. Motivo: {output_metadata.get('selection_reason', 'N/A')}"))

        candidates = output_metadata.get("candidate_stations", [])
        if candidates:
            st.caption(tr("INMET candidate ranking (best-first):", "Ranking de candidatos INMET (mejor primero):", "Ranking de candidatos INMET (melhor primeiro):"))
            st.dataframe(pd.DataFrame(candidates), use_container_width=True)
            st.caption(tr("Coverage ratio = fraction of expected hourly timestamps with records in requested period. Missing ratio = fraction of required variable cells that are missing over expected hourly grid.", "Coverage ratio = fracción de marcas horarias esperadas con registros en el periodo solicitado. Missing ratio = fracción de celdas de variables requeridas faltantes sobre la grilla horaria esperada.", "Coverage ratio = fração dos horários esperados com registros no período solicitado. Missing ratio = fração das células de variáveis obrigatórias ausentes na grade horária esperada."))
            st.caption(tr("Missing days columns indicate dates with missing INMET values for required variables and/or daily precipitation.", "Las columnas de días faltantes indican fechas con valores INMET faltantes para variables requeridas y/o precipitación diaria.", "As colunas de dias faltantes indicam datas com valores INMET ausentes para variáveis obrigatórias e/ou precipitação diária."))

        if not daily_storage and not hourly_records:
            status.update(
                label=tr(
                    "No data found for the selected inputs.",
                    "No se encontraron datos para las entradas seleccionadas.",
                    "Nenhum dado encontrado para as entradas selecionadas.",
                ),
                state="error",
                expanded=True,
            )
            st.error(
                tr(
                    "No weather records were returned. Try increasing INMET radius, adjusting date range, "
                    "or switching Source Selection to NASA only to test connectivity.",
                    "No se devolvieron registros meteorológicos. Prueba aumentar el radio INMET, ajustar el rango de fechas "
                    "o cambiar la selección de fuente a Solo NASA para validar conectividad.",
                    "Nenhum registro meteorológico foi retornado. Tente aumentar o raio INMET, ajustar o intervalo de datas "
                    "ou mudar a seleção da fonte para Somente NASA para validar conectividade.",
                )
            )
            st.caption(tr(f"Selection details: {output_metadata.get('selection_reason', 'N/A')}", f"Detalles de selección: {output_metadata.get('selection_reason', 'N/A')}", f"Detalhes da seleção: {output_metadata.get('selection_reason', 'N/A')}"))
            st.stop()

        # PHASE 3: WRITE OUT DATA
        ARM_COLS = ["Date", "Time", "Moisture Total", "Unit_1", "Precip", "Unit_2", "Irrigation", "Unit_3", "Type", "Type Description", "Interval", "Unit_4", "Leaf Wetness Duration", "Unit_5", "Min Temp", "Max Temp", "Avg Temp", "Temp Unit", "Min % Relative Humidity", "Max % Relative Humidity", "Avg % Relative Humidity", "Min Wind", "Max Wind", "Avg Wind", "Unit_6", "% Cloud Cover", "Avg Shortwave Radiation", "Unit_7", "Avg Soil Temp", "Unit_8", "0-10 cm Scaled Soil Moisture", "0-200 cm Scaled Soil Moisture", "Source", "Additional Comments"]
        ARM_DISPLAY = [c.split("_")[0] for c in ARM_COLS]

        if out_daily and daily_storage:
            st.write(tr("📊 Calculating Daily Statistics...", "📊 Calculando estadísticas diarias...", "📊 Calculando estatísticas diárias..."))
            sorted_dates = sorted(daily_storage.keys())
            
            if st.session_state.is_arm:
                arm_data = []
                for idx, dt in enumerate(sorted_dates):
                    d_map = daily_storage[dt]
                    prec = d_map.get("PRECTOTCORR")
                    prec_v = round(prec, 2) if prec is not None else None

                    t_vals, rh_vals, ws_vals = d_map.get("T2M", []), d_map.get("RH2M", []), d_map.get("WS2M", [])
                    t_min = round(min(t_vals), 2) if t_vals else None
                    t_max = round(max(t_vals), 2) if t_vals else None
                    t_avg = round(statistics.mean(t_vals), 2) if t_vals else None

                    rh_min = round(min(rh_vals), 2) if rh_vals else None
                    rh_max = round(max(rh_vals), 2) if rh_vals else None
                    rh_avg = round(statistics.mean(rh_vals), 2) if rh_vals else None

                    ws_min = round(min(ws_vals), 2) if ws_vals else None
                    ws_max = round(max(ws_vals), 2) if ws_vals else None
                    ws_avg = round(statistics.mean(ws_vals), 2) if ws_vals else None

                    # Convert wind speed from m/s to km/h for ARM output.
                    ws_min_kps = round(ws_min * 3.6, 2) if ws_min is not None else None
                    ws_max_kps = round(ws_max * 3.6, 2) if ws_max is not None else None
                    ws_avg_kps = round(ws_avg * 3.6, 2) if ws_avg is not None else None

                    arm_data.append([
                        to_arm_date(dt), "", prec_v, "mm" if prec_v is not None else "", prec_v, "mm" if prec_v is not None else "",
                        "", "", "RAIN" if prec and prec > 0 else "", "rain" if prec and prec > 0 else "", "", "", "", "",
                        t_min, t_max, t_avg, "C" if t_avg is not None else "", rh_min, rh_max, rh_avg, ws_min_kps, ws_max_kps, ws_avg_kps, "KPH" if ws_avg_kps is not None else "",
                        "", "", "", "", "", "", "", "ENTERED", ""
                    ])

                df = pd.DataFrame(arm_data, columns=ARM_DISPLAY)
                excel_daily_arm_buffer = io.BytesIO()
                with pd.ExcelWriter(excel_daily_arm_buffer, engine='openpyxl') as writer:
                    df.to_excel(writer, index=False, sheet_name="Meteorological_Data")
                    met_ws = writer.sheets["Meteorological_Data"]
                    # Columns holding decimal measurements (0-indexed into ARM_DISPLAY):
                    # 2/4=Moisture Total/Precip, 14-16=temps, 18-20=RH, 21-23=wind.
                    for col_idx in (2, 4, 14, 15, 16, 18, 19, 20, 21, 22, 23):
                        col_letter = get_column_letter(col_idx + 1)
                        for row in range(2, len(arm_data) + 2):
                            met_ws[f"{col_letter}{row}"].number_format = "0.00"

                    if enable_app_format:
                        # Build application worksheet in the same ARM workbook.
                        hourly_prec_by_dt = {}
                        for rec in hourly_records:
                            pval = rec.get("PRECTOTCORR")
                            if pval is None:
                                continue
                            try:
                                dt_local = datetime.strptime(f"{rec.get('date_key')} {int(float(rec.get('hr', 0))):02d}", "%Y%m%d %H")
                            except Exception:
                                continue
                            hourly_prec_by_dt[dt_local] = float(pval)

                        app_rows = []
                        one_decimal_rows = []
                        two_decimal_rows = []
                        for app_letter, app_date, app_time in app_dates_input:
                            app_dt = datetime.combine(app_date, app_time)

                            first_moisture_dt = None
                            for dtk in sorted(hourly_prec_by_dt.keys()):
                                if dtk >= app_dt and hourly_prec_by_dt[dtk] > 0:
                                    first_moisture_dt = dtk
                                    break

                            first_moisture_date = None
                            time_to_first = ""
                            time_unit = ""

                            if first_moisture_dt is not None:
                                first_moisture_date = first_moisture_dt.date()
                                delta_hrs = max(0, int(round((first_moisture_dt - app_dt).total_seconds() / 3600.0)))
                                if delta_hrs <= 24:
                                    time_to_first = delta_hrs
                                    time_unit = "HR"
                                else:
                                    time_to_first = max(1, int(delta_hrs // 24))
                                    time_unit = "DAY"
                            else:
                                cur_d = app_date
                                while cur_d <= fetch_end_date:
                                    dkey = cur_d.strftime("%Y%m%d")
                                    dprec = weather_result.daily_storage.get(dkey, {}).get("PRECTOTCORR")
                                    if dprec is not None and dprec > 0:
                                        first_moisture_date = cur_d
                                        time_to_first = (cur_d - app_date).days
                                        time_unit = "DAY"
                                        break
                                    cur_d += timedelta(days=1)

                            first_moisture_arm = to_arm_date(first_moisture_date.strftime("%Y%m%d")) if first_moisture_date else ""
                            first_moisture_amt = None
                            if first_moisture_date is not None:
                                dkey_fm = first_moisture_date.strftime("%Y%m%d")
                                dprec = weather_result.daily_storage.get(dkey_fm, {}).get("PRECTOTCORR")
                                if dprec is not None:
                                    first_moisture_amt = round(float(dprec), 1)

                            w2_before = get_precip_sum(app_date - timedelta(days=14), app_date - timedelta(days=8), weather_result.daily_storage)
                            w1_before = get_precip_sum(app_date - timedelta(days=7), app_date - timedelta(days=1), weather_result.daily_storage)
                            day_0 = get_precip_sum(app_date, app_date, weather_result.daily_storage)
                            h6_after = round(day_0 * 0.25, 2)
                            h24_after = day_0
                            w1_after = get_precip_sum(app_date + timedelta(days=1), app_date + timedelta(days=7), weather_result.daily_storage)
                            w2_after = get_precip_sum(app_date + timedelta(days=8), app_date + timedelta(days=14), weather_result.daily_storage)
                            w3_after = get_precip_sum(app_date + timedelta(days=15), app_date + timedelta(days=21), weather_result.daily_storage)
                            w4_after = get_precip_sum(app_date + timedelta(days=22), app_date + timedelta(days=28), weather_result.daily_storage)

                            block_start = len(app_rows)
                            app_rows.extend([
                                [f"--- Application {app_letter} ---", ""],
                                ["First Moisture Occured On", first_moisture_arm],
                                ["Time to First Moisture", time_to_first],
                                ["", time_unit],
                                ["Amount of First Moisture", first_moisture_amt],
                                ["", "mm"],
                                ["Moisture 2 Weeks Before Appl.", w2_before],
                                ["", "mm"],
                                ["Moisture 1 Week Before Appl.", w1_before],
                                ["", "mm"],
                                ["Moisture 6 Hours After Appl.", h6_after],
                                ["", "mm"],
                                ["Moisture 24 Hours After Appl.", h24_after],
                                ["", "mm"],
                                ["Moisture 1 Week After Appl.", w1_after],
                                ["", "mm"],
                                ["Moisture 2 Weeks After Appl.", w2_after],
                                ["", "mm"],
                                ["Moisture 3 Weeks After Appl.", w3_after],
                                ["", "mm"],
                                ["Moisture 4 Weeks After Appl.", w4_after],
                                ["", "mm"],
                                ["", ""],
                            ])
                            # Row offsets (within the block above) holding decimal mm values.
                            one_decimal_rows.append(block_start + 4)  # Amount of First Moisture
                            two_decimal_rows.extend(block_start + off for off in (6, 8, 10, 12, 14, 16, 18, 20))

                        df_apps = pd.DataFrame(app_rows, columns=["Field", "Value"])
                        df_apps.to_excel(writer, index=False, header=False, sheet_name="Weather_Application")
                        app_ws = writer.sheets["Weather_Application"]
                        for row_idx in one_decimal_rows:
                            app_ws.cell(row=row_idx + 1, column=2).number_format = "0.0"
                        for row_idx in two_decimal_rows:
                            app_ws.cell(row=row_idx + 1, column=2).number_format = "0.00"

                st.session_state.excel_daily_arm = excel_daily_arm_buffer.getvalue()
            
            else:
                f_daily = io.StringIO()
                writer_d = csv.writer(f_daily)
                d_header = ["DATE"]
                for p in hourly_req:
                    base = HEADER_MAP.get(p, p)
                    if p == "WS2M": d_header.extend([f"{base}_AVG", f"{base}_MAX", f"{base}_MIN"])
                    elif p == "WD2M": d_header.extend([f"{base}_AVG", HEADER_MAP.get("WD2M_COMPASS", "WD_CARDINAL")])
                    else: d_header.extend([f"{base}_AVG", f"{base}_MAX", f"{base}_MIN"])
                for p in daily_req: d_header.append(HEADER_MAP.get(p, p))
                writer_d.writerow(d_header)

                for dt in sorted_dates:
                    row = [dt]
                    for p in hourly_req:
                        vals = daily_storage[dt].get(p, [])
                        if not vals: row.extend(["", ""] if p == "WD2M" else ["", "", ""]); continue
                        if p == "WS2M" or p not in ["WD2M", "WS2M"]:
                            row.extend([f"{statistics.mean(vals):.2f}", f"{max(vals):.2f}", f"{min(vals):.2f}"])
                        elif p == "WD2M":
                            row.extend([f"{vector_average_degrees(vals):.2f}", deg_to_compass_16(vector_average_degrees(vals))])
                    for p in daily_req:
                        val = daily_storage[dt].get(p)
                        row.append(f"{val:.2f}" if val is not None and val != -999 else "")
                    writer_d.writerow(row)
                st.session_state.csv_daily_str = f_daily.getvalue()

        if out_hourly and not st.session_state.is_arm and hourly_records:
            f_hourly = io.StringIO()
            writer_h = csv.writer(f_hourly)
            out_h = ["DATE", "HR"]
            for p in hourly_req:
                out_h.append(HEADER_MAP.get(p, p))
            include_hourly_precip = output_metadata.get("primary_source") == "INMET"
            if include_hourly_precip:
                out_h.append("PREC_MM_HR")
            if "WD2M" in hourly_req:
                out_h.append(HEADER_MAP.get("WD2M_COMPASS", "WD_CARDINAL"))
            writer_h.writerow(out_h)

            for rec in hourly_records:
                row_vals = [rec.get("date_key", ""), rec.get("hr", "")]
                for p in hourly_req:
                    v = rec.get(p)
                    row_vals.append(v if v is not None else "")
                if include_hourly_precip:
                    pval = rec.get("PRECTOTCORR")
                    row_vals.append(pval if pval is not None else "")
                if "WD2M" in hourly_req:
                    wd_val = rec.get("WD2M")
                    row_vals.append(deg_to_compass_16(wd_val) if wd_val is not None else "")
                writer_h.writerow(row_vals)

            st.session_state.csv_hourly_str = f_hourly.getvalue()

        if out_hourly and st.session_state.is_arm and hourly_records:
            arm_hr_data = []
            for idx, r in enumerate(hourly_records):
                t_val = round(r['T2M'], 2) if r.get('T2M') is not None else None
                rh_val = round(r['RH2M'], 2) if r.get('RH2M') is not None else None
                ws_val = round(r['WS2M'] * 3.6, 2) if r.get('WS2M') is not None else None
                time_str = f"{str(r['hr']).split('.')[0].zfill(2)}:00"
                arm_hr_data.append([
                    to_arm_date(r['date_key']), time_str, "", "", "", "", "", "", "", "", "", "", "", "",
                    "", t_val, "C" if t_val is not None else "", "", "", rh_val, "", "", ws_val, "KPH" if ws_val is not None else "",
                    "", "", "", "", "", "", "", "", "ENTERED", ""
                ])
            df_hr = pd.DataFrame(arm_hr_data, columns=ARM_DISPLAY)
            excel_hourly_arm_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_hourly_arm_buffer, engine='openpyxl') as writer:
                df_hr.to_excel(writer, index=False)
                hr_ws = writer.sheets["Sheet1"]
                # Columns actually holding t_val/rh_val/ws_val above (0-indexed into ARM_DISPLAY): 15, 19, 22.
                for col_idx in (15, 19, 22):
                    col_letter = get_column_letter(col_idx + 1)
                    for row in range(2, len(arm_hr_data) + 2):
                        hr_ws[f"{col_letter}{row}"].number_format = "0.00"
            st.session_state.excel_hourly_arm = excel_hourly_arm_buffer.getvalue()

        if enable_app_format and daily_storage and not st.session_state.is_arm:
            st.write(tr("🌱 Generating Application Layout...", "🌱 Generando formato de aplicaciones...", "🌱 Gerando formato de aplicações..."))
            app_table_data = []
            for app_letter, app_date, _app_time in app_dates_input:
                w2_before = get_precip_sum(app_date - timedelta(days=14), app_date - timedelta(days=8), weather_result.daily_storage)
                w1_before = get_precip_sum(app_date - timedelta(days=7), app_date - timedelta(days=1), weather_result.daily_storage)
                day_0 = get_precip_sum(app_date, app_date, weather_result.daily_storage)
                h6_after = round(day_0 * 0.25, 2)
                h24_after = day_0
                w1_after = get_precip_sum(app_date + timedelta(days=1), app_date + timedelta(days=7), weather_result.daily_storage)
                w2_after = get_precip_sum(app_date + timedelta(days=8), app_date + timedelta(days=14), weather_result.daily_storage)
                w3_after = get_precip_sum(app_date + timedelta(days=15), app_date + timedelta(days=21), weather_result.daily_storage)
                w4_after = get_precip_sum(app_date + timedelta(days=22), app_date + timedelta(days=28), weather_result.daily_storage)

                app_table_data.extend([
                    [f"--- Application {app_letter} ---", "", ""],
                    ["Moisture 2 Weeks Before Appl.", w2_before, "mm"],
                    ["Moisture 1 Week Before Appl.", w1_before, "mm"],
                    ["Moisture 6 Hours After Appl.", h6_after, "mm"],
                    ["Moisture 24 Hours After Appl.", h24_after, "mm"],
                    ["Moisture 1 Week After Appl.", w1_after, "mm"],
                    ["Moisture 2 Weeks After Appl.", w2_after, "mm"],
                    ["Moisture 3 Weeks After Appl.", w3_after, "mm"],
                    ["Moisture 4 Weeks After Appl.", w4_after, "mm"],
                    ["", "", ""] 
                ])

            df_apps = pd.DataFrame(app_table_data, columns=["Interval", "Value", "Unit"])
            excel_app_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_app_buffer, engine='openpyxl') as writer: df_apps.to_excel(writer, index=False, header=False, sheet_name="Application_Moisture")
            st.session_state.excel_app_format = excel_app_buffer.getvalue()

        status.update(label=tr("Data Processing Complete!", "Procesamiento de datos completado!", "Processamento de dados concluído!"), state="complete", expanded=False)
        if debug_mode and debug_log: st.session_state.debug_log = debug_log

# --- Rendering Persisted Download Buttons ---
# This block runs independently of the "DOWNLOAD & PROCESS" button so it won't disappear on click
if any([st.session_state.csv_hourly_str, st.session_state.csv_daily_str, st.session_state.excel_hourly_arm, st.session_state.excel_daily_arm, st.session_state.excel_app_format]):
    st.divider()
    st.success(tr("✅ Downloads are ready!", "✅ Descargas listas!", "✅ Downloads prontos!"))
    
    cols_dl = st.columns(3)
    idx_dl = 0
    
    if st.session_state.excel_hourly_arm:
        with cols_dl[idx_dl % 3]: st.download_button(tr("⬇️ Download Hourly Data (Excel)", "⬇️ Descargar datos horarios (Excel)", "⬇️ Baixar dados horários (Excel)"), data=st.session_state.excel_hourly_arm, file_name=f"{st.session_state.base_filename}_Hourly_ARM.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
        idx_dl += 1
    elif st.session_state.csv_hourly_str:
        with cols_dl[idx_dl % 3]: st.download_button(tr("⬇️ Download Hourly Data (CSV)", "⬇️ Descargar datos horarios (CSV)", "⬇️ Baixar dados horários (CSV)"), data=st.session_state.csv_hourly_str, file_name=f"{st.session_state.base_filename}_Hourly.csv", mime="text/csv", use_container_width=True)
        idx_dl += 1

    if st.session_state.excel_daily_arm:
        with cols_dl[idx_dl % 3]: st.download_button(tr("⬇️ Download Daily Stats (Excel)", "⬇️ Descargar estadísticas diarias (Excel)", "⬇️ Baixar estatísticas diárias (Excel)"), data=st.session_state.excel_daily_arm, file_name=f"{st.session_state.base_filename}_DailyStats_ARM.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
        idx_dl += 1
    elif st.session_state.csv_daily_str:
        with cols_dl[idx_dl % 3]: st.download_button(tr("⬇️ Download Daily Stats (CSV)", "⬇️ Descargar estadísticas diarias (CSV)", "⬇️ Baixar estatísticas diárias (CSV)"), data=st.session_state.csv_daily_str, file_name=f"{st.session_state.base_filename}_DailyStats.csv", mime="text/csv", use_container_width=True)
        idx_dl += 1

    if st.session_state.excel_app_format:
        with cols_dl[idx_dl % 3]:
            st.download_button(tr("⬇️ Download Application Format", "⬇️ Descargar formato de aplicación", "⬇️ Baixar formato de aplicação"), data=st.session_state.excel_app_format, file_name=f"{st.session_state.base_filename}_AppFormat.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
    if st.session_state.output_metadata_json:
        with cols_dl[idx_dl % 3]:
            st.download_button(tr("⬇️ Download Source Metadata (JSON)", "⬇️ Descargar metadatos de fuente (JSON)", "⬇️ Baixar metadados da fonte (JSON)"), data=st.session_state.output_metadata_json, file_name=f"{st.session_state.base_filename}_metadata.json", mime="application/json", use_container_width=True)

    if hasattr(st.session_state, "debug_log") and st.session_state.debug_log:
        with st.expander(tr("Show Debug Logs", "Mostrar logs de depuración", "Mostrar logs de depuração")): st.code("\n".join(st.session_state.debug_log))
